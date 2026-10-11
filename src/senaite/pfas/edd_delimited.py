# -*- coding: utf-8 -*-
"""The column-mapped delimited-file EDD: one generic writer for any
delivery format whose files, columns, codes and checks are a profile's DATA.

A profile of this type (format "delimited") holds:

    files     [{name, grain, delimiter, header, header_prefix, newline,
                columns: [{heading, source, fallback, constant, map, format,
                           when, join, required, list, max_length}]}]
    maps      {name: {by, list, rows: [{ours, code, description}]}}
    lists     {name: [code, ...]}            the format's valid values
    defaults  {name: value}                  the lab's identity values
    client_fields [{key, label, list}]       what a client gives
    package   the archive's name pattern when there is more than one file
    analyte_naming  (as every profile: the format's names and code overrides)

The writer reads FACTS (edd_facts collects them from the platform; tests
build them by hand):

    {"batch": {id, sdg, title},
     "client": {field: value},
     "samples": [{<sample facts>, "results": [{<result facts>, "qc": {...}}]}]}

Grains: "sample" (one row per sample), "result" (one row per result),
"test_batch" (one row per sample test and lab batch: PREP for the
extraction, ANALYSIS for the run). A column carries a source (a fact), a
constant, or nothing; "when" writes it only if a condition holds; a map
turns the value into the format's code (a value the map lacks is refused,
never passed through); a format writes a date, a CAS or a case; required,
list and max_length are checked, and a breach blocks the delivery.

Pure; Python 2.7 and 3.
"""
from __future__ import absolute_import, unicode_literals

import io
import re
import zipfile
from datetime import date, datetime

try:                                         # Python 2
    text_type = unicode                      # noqa: F821
except NameError:                            # Python 3
    text_type = str

GRAINS = (("sample", u"One row per sample"),
          ("result", u"One row per result"),
          ("test_batch", u"One row per sample test and lab batch (PREP, ANALYSIS)"))
DELIMITERS = ((",", u"Comma"), ("tab", u"Tab"), (";", u"Semicolon"), ("|", u"Pipe"))
HEADERS = (("names", u"Column headings"), ("comment", u"Column headings as a comment row"),
           ("none", u"No header row"))
NEWLINES = (("crlf", u"CRLF (Windows)"), ("lf", u"LF"))

# the platform facts a column can carry: (source, label)
SOURCES = (
    ("sample.id", u"Sample: lab sample ID"),
    ("sample.client_sample_id", u"Sample: client sample ID"),
    ("sample.date_sampled", u"Sample: date sampled"),
    ("sample.date_received", u"Sample: date received"),
    ("sample.date_prepared", u"Sample: date prepared (a lab QC sample: its extraction's finish)"),
    ("sample.matrix", u"Sample: matrix (sample type)"),
    ("sample.qc_type", u"Sample: QC type (ours; NA for a field sample)"),
    ("sample.origin", u"Sample: field or lab"),
    ("sample.sample_point", u"Sample: sample point"),
    ("sample.coc_number", u"Sample: CoC number"),
    ("sample.duplicate_of", u"Sample: the sample it duplicates or spikes (its sample ID)"),
    ("batch.id", u"Batch: ID"),
    ("batch.sdg", u"Batch: sample delivery group"),
    ("batch.title", u"Batch: title"),
    ("analysis.keyword", u"Analyte: keyword"),
    ("analysis.cas", u"Analyte: CAS (the profile's code override wins)"),
    ("analysis.name", u"Analyte: name in this format (Analytes tab)"),
    ("analysis.role", u"Analyte: role (target, tentative, surrogate, internal standard)"),
    ("analysis.method", u"Analysis: method"),
    ("analysis.test_kind", u"Analysis: test (initial, dilution 1, ...)"),
    ("analysis.acquired_at", u"Analysis: analysed (acquired)"),
    ("analysis.run_date", u"Analysis: the run's date (lab QC)"),
    ("analysis.judged_at", u"Analysis: when the worker judged the run's QC (lab QC)"),
    ("analysis.worksheet", u"Analysis: worksheet (extraction batch)"),
    ("result.value", u"Result: value as reported"),
    ("result.detected", u"Result: detected / not detected"),
    ("result.unit", u"Result: unit"),
    ("result.rl", u"Result: reporting limit"),
    ("result.mdl", u"Result: method detection limit"),
    ("result.dilution_factor", u"Result: dilution factor"),
    ("result.qualifiers", u"Result: qualifiers (each mapped, joined)"),
    ("result.reportable", u"Result: reported / superseded"),
    ("result.prep_date", u"Result: prepared (extraction finished)"),
    ("qc.measured", u"QC: measured concentration"),
    ("qc.recovery", u"QC: recovery (%)"),
    ("qc.recovery_low", u"QC: recovery lower limit (%)"),
    ("qc.recovery_high", u"QC: recovery upper limit (%)"),
    ("qc.recovery_outside", u"QC: recovery outside its limits (outside / inside)"),
    ("qc.rpd", u"QC: RPD (%)"),
    ("qc.rpd_limit", u"QC: RPD limit (%)"),
    ("qc.rpd_outside", u"QC: RPD outside its limit (outside / inside)"),
    ("row.batch_kind", u"Lab batch: kind (PREP, ANALYSIS)"),
    ("row.batch_id", u"Lab batch: ID"),
)
SOURCE_KEYS = frozenset(s for s, _l in SOURCES)
# a source naming a client field or a profile default: "client.<key>",
# "profile.<name>"
DYNAMIC = ("client.", "profile.")

_CAS = re.compile(r"^(\d{2,7})-?(\d{2})-?(\d)$")
_COND = re.compile(r"^\s*(not\s+)?([a-z_]+\.[A-Za-z0-9_]+)\s*(?:(=|!=|in|not in)\s*(.*?))?\s*$")


def known_source(source):
    return source in SOURCE_KEYS or any(source.startswith(p) and len(source) > len(p) for p in DYNAMIC)


# ── a profile's own problems (shown on its pages; a delivery with any is
# refused) ───────────────────────────────────────────────────────────────────

def parse_condition(text):
    """(negate, source, op, values) or None when it cannot be read.
    Forms: "src", "not src", "src = v", "src != v", "src in a,b",
    "src not in a,b"."""
    m = _COND.match(text or u"")
    if not m:
        return None
    neg, src, op, rest = m.groups()
    if op and neg:
        return None
    if op and (rest is None or rest == u""):
        return None
    values = [v.strip() for v in (rest or u"").split(u",")] if op in ("in", "not in") else (
        [rest.strip()] if op else [])
    return (bool(neg), src, op or u"", values)


def problems(profile):
    """["what is wrong", ...] with a profile's files, columns and maps."""
    p = profile or {}
    out = []
    files = p.get("files") or []
    if not files:
        out.append(u"No file is set up (Files tab).")
    names = [f.get("name") or u"" for f in files]
    for n in sorted(set(x for x in names if names.count(x) > 1)):
        out.append(u"Two files are named %s." % n)
    maps, lists = p.get("maps") or {}, p.get("lists") or {}
    grains = dict(GRAINS)
    for f in files:
        fname = f.get("name") or u"(a file without a name)"
        if not f.get("name"):
            out.append(u"A file has no name.")
        if f.get("grain") not in grains:
            out.append(u"%s: choose its rows (one per sample, result, or test and lab batch)." % fname)
        if not f.get("columns"):
            out.append(u"%s has no columns." % fname)
        for i, c in enumerate(f.get("columns") or [], 1):
            where = u"%s column %d (%s)" % (fname, i, c.get("heading") or u"no heading")
            if not (c.get("heading") or u"").strip():
                out.append(u"%s: no heading." % where)
            for key in ("source", "fallback"):
                s = c.get(key) or u""
                if s and not known_source(s):
                    out.append(u"%s: unknown %s %s." % (where, key, s))
            if c.get("map") and c["map"] not in maps:
                out.append(u"%s: no code map named %s." % (where, c["map"]))
            if c.get("list") and c["list"] not in lists:
                out.append(u"%s: no value list named %s." % (where, c["list"]))
            if c.get("when"):
                cond = parse_condition(c["when"])
                if cond is None or not known_source(cond[1]):
                    out.append(u"%s: the condition %r cannot be read." % (where, c["when"]))
            fmt = c.get("format") or u""
            if fmt and not (fmt.startswith(u"date:") or fmt in (u"cas:dashed", u"cas:digits", u"upper",
                                                                  u"lower")):
                out.append(u"%s: unknown format %s." % (where, fmt))
            if fmt.startswith(u"date:"):
                try:
                    datetime(2026, 1, 2, 3, 4).strftime(str(fmt[5:]))
                except (ValueError, TypeError):
                    out.append(u"%s: the date pattern %s cannot be used." % (where, fmt[5:]))
            ml = c.get("max_length")
            if ml not in (None, u"") and not (u"%s" % ml).isdigit():
                out.append(u"%s: the maximum length must be a number." % where)
    for name, m in sorted(maps.items()):
        if m.get("by") and not known_source(m["by"]):
            out.append(u"Code map %s: unknown second key %s." % (name, m["by"]))
        if m.get("list") and m["list"] not in lists:
            out.append(u"Code map %s: no value list named %s." % (name, m["list"]))
        ours = [r.get("ours") for r in m.get("rows") or []]
        for o in sorted(set(x for x in ours if ours.count(x) > 1)):
            out.append(u"Code map %s: %s is mapped twice." % (name, o))
        allowed = set(lists.get(m.get("list")) or []) if m.get("list") else None
        for r in m.get("rows") or []:
            if allowed is not None and r.get("code") and r["code"] not in allowed:
                out.append(u"Code map %s: %s is not in the list %s." % (name, r["code"], m["list"]))
    return out


# ── rows ────────────────────────────────────────────────────────────────────

def _flat(prefix, d, out):
    for k, v in (d or {}).items():
        if k in ("results", "qc"):
            continue
        out[u"%s.%s" % (prefix, k)] = v


def _context(facts, profile, sample=None, result=None):
    ctx = {}
    _flat(u"batch", facts.get("batch"), ctx)
    _flat(u"client", facts.get("client"), ctx)
    _flat(u"profile", (profile or {}).get("defaults"), ctx)
    if sample is not None:
        _flat(u"sample", sample, ctx)
    if result is not None:
        for k in ("keyword", "cas", "name", "role", "method", "test_kind", "acquired_at", "run_date",
                  "judged_at", "worksheet"):
            ctx[u"analysis.%s" % k] = result.get(k)
        for k in ("value", "detected", "unit", "rl", "mdl", "dilution_factor", "qualifiers", "reportable",
                  "prep_date"):
            ctx[u"result.%s" % k] = result.get(k)
        _flat(u"qc", result.get("qc"), ctx)
    return ctx


def _results(sample, profile):
    """The sample's results with the profile's analyte names and code
    overrides applied."""
    naming = (profile or {}).get("analyte_naming") or {}
    out = []
    for r in sample.get("results") or []:
        r = dict(r)
        nm = naming.get(r.get("keyword")) or {}
        if nm.get("code_override"):
            r["cas"] = nm["code_override"]
        r["name"] = nm.get("parameter_name") or u""
        if nm.get("tic") and r.get("role") == u"target":
            r["role"] = u"tentative"
        out.append(r)
    return out


def rows_for(grain, facts, profile):
    """[(row label, context)] for one file."""
    out = []
    for s in facts.get("samples") or []:
        label = s.get("id") or u"?"
        if grain == "sample":
            out.append((label, _context(facts, profile, s)))
            continue
        results = _results(s, profile)
        if grain == "result":
            for r in results:
                out.append((u"%s %s" % (label, r.get("keyword")), _context(facts, profile, s, r)))
        elif grain == "test_batch":
            seen = set()
            for r in results:
                key = (r.get("method"), r.get("acquired_at") or r.get("judged_at") or r.get("run_date"),
                       r.get("test_kind"))
                if key in seen:
                    continue
                seen.add(key)
                kinds = []
                if r.get("prep_date") and r.get("worksheet"):
                    kinds.append((u"PREP", r.get("worksheet")))
                if r.get("worksheet"):
                    kinds.append((u"ANALYSIS", r.get("worksheet")))
                for kind, bid in kinds:
                    ctx = _context(facts, profile, s, r)
                    ctx[u"row.batch_kind"], ctx[u"row.batch_id"] = kind, bid
                    out.append((u"%s %s %s" % (label, r.get("test_kind") or u"", kind), ctx))
    return out


# ── cells ───────────────────────────────────────────────────────────────────

def _blank(v):
    return v is None or (isinstance(v, (list, tuple)) and not v) or (u"%s" % v).strip() == u""


def holds(cond_text, ctx):
    cond = parse_condition(cond_text)
    if cond is None:
        return False
    neg, src, op, values = cond
    v = ctx.get(src)
    if not op:
        ok = not _blank(v)
        return not ok if neg else ok
    text = u"" if v is None else u"%s" % v
    if op == "=":
        return text == values[0]
    if op == "!=":
        return text != values[0]
    if op == "in":
        return text in values
    return text not in values                                   # not in


def _as_datetime(v):
    """A date fact given as ISO text (facts read from JSON) as a datetime."""
    if isinstance(v, (datetime, date)) or not v:
        return v
    for pat in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime((u"%s" % v)[:len(datetime(2026, 1, 1).strftime(pat))], pat)
        except ValueError:
            continue
    return v


def _date_text(v, pattern):
    v = _as_datetime(v)
    if isinstance(v, (datetime, date)):
        try:
            return text_type(v.strftime(str(pattern)))
        except ValueError:
            return u""
    return u"" if v is None else u"%s" % v


def _cas(value, how):
    m = _CAS.match(value or u"")
    if not m:
        return value                         # a format's own code is written as entered
    return u"%s-%s-%s" % m.groups() if how == "dashed" else u"".join(m.groups())


def _number(v):
    if isinstance(v, float):
        return (u"%r" % v) if v != int(v) else (u"%d" % int(v))
    return u"%s" % v


def lookup(maps, name, value, ctx):
    """The map's code for a value (keyed by the map's second source first,
    "value|second"), or None when the map has none."""
    m = maps.get(name) or {}
    rows = dict(((r.get("ours") or u""), r.get("code") or u"") for r in m.get("rows") or [])
    v = u"%s" % value
    if m.get("by"):
        second = ctx.get(m["by"])
        key = u"%s|%s" % (v, u"" if second is None else second)
        if key in rows:
            return rows[key]
    return rows.get(v)


def cell(col, ctx, profile, errors, where):
    """The text one column writes for one row (errors appended)."""
    maps = (profile or {}).get("maps") or {}
    if col.get("when") and not holds(col["when"], ctx):
        value = None
    elif not _blank(col.get("constant")):
        value = col["constant"]
    elif col.get("source"):
        value = ctx.get(col["source"])
        if _blank(value) and col.get("fallback"):
            value = ctx.get(col["fallback"])
    else:
        value = None
    fmt = col.get("format") or u""
    if fmt.startswith(u"date:"):
        value = _date_text(value, fmt[5:]) if not isinstance(value, (list, tuple)) else value
    values = list(value) if isinstance(value, (list, tuple)) else ([] if _blank(value) else [value])
    if col.get("map"):
        coded = []
        for v in values:
            code = lookup(maps, col["map"], v, ctx)
            if code is None:
                errors.append({"row": where, "field": col.get("heading"), "type": "BLOCKING",
                               "message": u"%s: %r has no code in the map %s -- add it (Code maps)"
                                          % (col.get("heading"), u"%s" % v, col["map"])})
                continue
            if code and code not in coded:
                coded.append(code)
        values = coded
    text = (col.get("join") or u"").join(
        _date_text(v, u"%Y-%m-%d %H:%M") if isinstance(v, (datetime, date)) else _number(v) for v in values)
    if fmt.startswith(u"cas:"):
        text = _cas(text, fmt[4:])
    elif fmt == u"upper":
        text = text.upper()
    elif fmt == u"lower":
        text = text.lower()
    if col.get("required") and text.strip() == u"":
        errors.append({"row": where, "field": col.get("heading"), "type": "BLOCKING",
                       "message": u"%s is required and blank" % col.get("heading")})
    lst = col.get("list")
    if lst and text != u"" and text not in ((profile or {}).get("lists") or {}).get(lst, ()):
        errors.append({"row": where, "field": col.get("heading"), "type": "BLOCKING",
                       "message": u"%s: %r is not in the value list %s" % (col.get("heading"), text, lst)})
    ml = col.get("max_length")
    if ml not in (None, u"") and (u"%s" % ml).isdigit() and len(text) > int(ml):
        errors.append({"row": where, "field": col.get("heading"), "type": "BLOCKING",
                       "message": u"%s: %r is longer than %s characters" % (col.get("heading"), text, ml)})
    return text


# ── files ───────────────────────────────────────────────────────────────────

def _quote(text, delim):
    if any(ch in text for ch in (delim, u'"', u"\r", u"\n")):
        return u'"%s"' % text.replace(u'"', u'""')
    return text


def render(f, rows):
    """The file's text: header as set, then each row."""
    delim = f.get("delimiter") or u","
    if delim == u"tab":
        delim = u"\t"
    nl = u"\n" if f.get("newline") == "lf" else u"\r\n"
    lines = []
    headings = [c.get("heading") or u"" for c in f.get("columns") or []]
    header = f.get("header") or "names"
    if header == "names":
        lines.append(delim.join(_quote(h, delim) for h in headings))
    elif header == "comment" and headings:
        prefix = f.get("header_prefix") or u"#"
        lines.append(delim.join(_quote(h, delim) for h in [prefix + headings[0]] + headings[1:]))
    for r in rows:
        lines.append(delim.join(_quote(v, delim) for v in r))
    return nl.join(lines) + (nl if lines else u"")


def _name(pattern, tokens, errors):
    try:
        return (pattern or u"").format(**tokens)
    except (KeyError, IndexError, ValueError) as exc:
        errors.append({"row": 0, "field": "name", "type": "BLOCKING",
                       "message": u"the name pattern %r cannot be filled (%s)" % (pattern, exc)})
        return pattern or u"EDD"


def write(profile, facts, today=None):
    """{"files": [(name, bytes)], "errors": [...], "rows": {name: n},
    "delivery": (name, bytes)} -- the delivery is the one file, or a zip of
    them all."""
    errors = [{"row": 0, "field": "profile", "type": "BLOCKING", "message": m} for m in problems(profile)]
    today = today or date.today()
    batch = facts.get("batch") or {}
    tokens = {"batch_id": batch.get("id") or u"", "batch": batch.get("id") or u"",
              "sdg": batch.get("sdg") or batch.get("id") or u"",
              "date": today.strftime("%Y%m%d")}
    files, counts = [], {}
    if not any(e["type"] == "BLOCKING" for e in errors):
        for f in (profile or {}).get("files") or []:
            name = _name(f.get("name"), tokens, errors)
            out_rows = []
            for label, ctx in rows_for(f.get("grain"), facts, profile):
                out_rows.append([cell(c, ctx, profile, errors, u"%s: %s" % (name, label))
                                 for c in f.get("columns") or []])
            if not out_rows:
                errors.append({"row": 0, "field": name, "type": "BLOCKING",
                               "message": u"%s has no rows: the batch holds nothing it reports" % name})
            files.append((name, render(f, out_rows).encode("utf-8")))
            counts[name] = len(out_rows)
    delivery = None
    if len(files) == 1:
        delivery = files[0]
    elif files:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for name, data in files:
                info = zipfile.ZipInfo(str(name), date_time=(today.year, today.month, today.day, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o644 << 16        # a plain file when unzipped, not mode 000
                z.writestr(info, data)
        delivery = (_name((profile or {}).get("package") or u"{sdg}_{date}_EDD.zip", tokens, errors),
                    buf.getvalue())
    return {"files": files, "errors": errors, "rows": counts, "delivery": delivery}


# ── value lists from the agency's workbook ──────────────────────────────────

def read_list(xlsx_bytes, sheet, column):
    """The values of one column of one sheet (its header row names the
    column; a leading '#' on the header is ignored). (values, error)."""
    try:
        import openpyxl
    except ImportError:
        return [], u"openpyxl is required to read a workbook"
    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    except Exception as exc:                                     # noqa: BLE001
        return [], u"the workbook cannot be opened (%s)" % exc
    if sheet not in wb.sheetnames:
        return [], u"the workbook has no sheet %s" % sheet
    rows = wb[sheet].iter_rows(values_only=True)
    try:
        header = [(u"%s" % h).strip().lstrip(u"#") if h is not None else u"" for h in next(rows)]
    except StopIteration:
        return [], u"the sheet %s is empty" % sheet
    if column not in header:
        return [], u"the sheet %s has no column %s" % (sheet, column)
    i = header.index(column)
    out = []
    for r in rows:
        v = r[i] if i < len(r) else None
        if v is not None and (u"%s" % v).strip() and (u"%s" % v).strip() not in out:
            out.append((u"%s" % v).strip())
    return out, u""


def skeleton(name, state=u""):
    """A new profile of this type: one file, no columns yet."""
    return {"name": name, "format": "delimited", "base": "delimited", "state": state,
            "files": [{"name": u"{sdg}_{date}.csv", "grain": "result", "delimiter": u",",
                       "header": "names", "header_prefix": u"#", "newline": "crlf", "columns": []}],
            "maps": {}, "lists": {}, "defaults": {}, "client_fields": [],
            "package": u"{sdg}_{date}_EDD.zip", "analyte_naming": {},
            "notes": u"Column-mapped delimited files."}
